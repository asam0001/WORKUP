import os
import sys

# Ensure python-pptx site-packages is in path
base_pkg = r"C:\Users\Mruthyunjaya\AppData\Local\uv\cache\archive-v0\OAAt0ANKOFoe2tgg\Lib\site-packages"
if base_pkg not in sys.path:
    sys.path.insert(0, base_pkg)

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Color Palette
COLOR_NAVY = RGBColor(15, 41, 66)        # #0F2942 - Primary Header
COLOR_SLATE = RGBColor(30, 41, 59)       # #1E293B - Body Text
COLOR_BLUE = RGBColor(30, 95, 138)       # #1E5F8A - Secondary Accent
COLOR_AZURE = RGBColor(37, 99, 235)      # #2563EB - Highlights & Badges
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B - Subtitles & Footers
COLOR_LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC - Slide Background
COLOR_CARD_BG = RGBColor(255, 255, 255)  # #FFFFFF - Card Background
COLOR_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 - Card Border
COLOR_GREEN = RGBColor(5, 150, 105)      # #059669 - Success / Completed
COLOR_AMBER = RGBColor(217, 119, 6)      # #D97706 - Warning / In-Progress
COLOR_RED = RGBColor(220, 38, 38)        # #DC2626 - Alert / Deficit

TOTAL_SLIDES = 23

def set_slide_background(slide, color=COLOR_LIGHT_BG):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, tracker, title, subtitle=None):
    # Left accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.38), Inches(0.08), Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_AZURE
    bar.line.fill.background()

    # Title text box
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(0.32), Inches(8.4), Inches(0.95))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = tracker.upper()
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_AZURE

    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY

    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_MUTED

    # Top right pill: Institution
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(0.42), Inches(3.03), Inches(0.42))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(239, 246, 255)
    pill.line.color.rgb = RGBColor(191, 219, 254)
    pill.line.width = Pt(1)
    ptf = pill.text_frame
    ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]
    pp.text = "ACE Engineering College | CSE (AI & DS)"
    pp.alignment = PP_ALIGN.CENTER
    pp.font.size = Pt(8.5)
    pp.font.bold = True
    pp.font.color.rgb = COLOR_BLUE

def add_footer(slide, current_slide):
    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_CARD_BORDER
    line.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    
    r1 = p.add_run()
    r1.text = "Academic Year 2026–2027  |  Major Project Review – Phase 1"
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = COLOR_MUTED

    r_mid = p.add_run()
    r_mid.text = " " * 32 + "WORKUP: Virtual Professional Experience Platform" + " " * 32
    r_mid.font.size = Pt(8.5)
    r_mid.font.bold = True
    r_mid.font.color.rgb = COLOR_NAVY

    r2 = p.add_run()
    r2.text = f"Slide {current_slide} of {TOTAL_SLIDES}"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = COLOR_MUTED

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

def add_badge(slide, left, top, width, height, text, bg_color, text_color):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    b.fill.solid()
    b.fill.fore_color.rgb = bg_color
    b.line.fill.background()
    tf = b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = text_color
    return b
